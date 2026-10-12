#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Qinhuangdao DAS raw .dat -> 1-min RMS Excel
============================================

目标
----
从原始 DAS 200 Hz .dat 文件中，仅处理指定时间段：
    2025-06-16 13:30:00 -> 2025-06-21 09:30:00

原始目录结构示例：
    你选择的总文件夹/
    ├─ 2025-06-16/
    │  ├─ 2025-06-16-13-29-58-out.dat
    │  ├─ 2025-06-16-13-30-03-out.dat
    │  └─ ...
    ├─ 2025-06-17/
    ├─ 2025-06-18/
    ├─ ...
    └─ 2025-06-21/

.dat 格式依据你之前的读取代码：
    - 前 64 个 float32 为 header
    - C[10] = sampling frequency
    - C[16] = number of channels
    - C[17] = file length in seconds
    - payload reshape 为 (N_channels, N_samples), order='F'

输出
----
自动在所选总文件夹下建立：
    DAS_1min_RMS_20250616_133000_to_20250621_093000/

其中：
    Coast_1min_RMS.xlsx
    Lagoon1_1min_RMS.xlsx
    Lagoon2_1min_RMS.xlsx
    Lagoon3_1min_RMS.xlsx
    minute_coverage.csv
    processing_summary.txt

Excel 结构与之前 30-min RMS 文件保持一致：
    - 第一列：channel
    - 第一行：1-min 时间
    - 行：channel
    - 列：time
    - 单元格：该分钟内的 demeaned RMS

内存设计
--------
1. 不拼接 5 天原始数据；
2. 每次只 mmap 一个 5 s 文件；
3. 只读取研究所需通道；
4. 逐文件累加 sum / sumsq / count；
5. 累加器与最终 RMS 使用磁盘 memmap，而不是全部放 RAM；
6. Excel 使用 openpyxl write-only 流式写出。

注意
----
- 当前时间区间采用 [START_TIME, END_TIME)，即不包含 END_TIME 本身。
  因此最后一个 1-min 时间标签是 2025-06-21 09:29:00。
- 这一阶段只做 1-min RMS，不做 log10、winsorization、detrend、标准化或频谱分析。
- 默认分钟有效覆盖率 < 90% 时，该分钟设为 NaN。
"""

from __future__ import annotations

import gc
import math
import os
import re
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox
except Exception as exc:
    raise RuntimeError("无法导入 tkinter。请使用带 Tk 支持的 Python 运行。") from exc

try:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
except Exception as exc:
    raise RuntimeError(
        "缺少 openpyxl。请先运行：pip install openpyxl"
    ) from exc


# =============================================================================
# 1. 用户配置
# =============================================================================

START_TIME = datetime(2025, 6, 16, 13, 30, 0)
END_TIME   = datetime(2025, 6, 21, 9, 30, 0)  # exclusive

# 你的四个 DAS 区段，均按 Python 原代码中的通道索引理解。
# 例如原代码 D1[413, :] 对应这里的 channel 413。
SITE_CHANNEL_RANGES: Dict[str, Tuple[int, int]] = {
    "Coast":    (330, 430),
    "Lagoon1":  (32, 101),
    "Lagoon2":  (102, 163),
    "Lagoon3":  (164, 257),
}

# 一分钟至少达到多少原始样本覆盖率才保留。
MIN_MINUTE_COVERAGE = 0.90

# 文件名时间是否表示“该 5 s 文件的起始时间”。
FILE_TIMESTAMP_IS_START = True

# 默认只输出四个研究区段，避免重复写一个巨大 all-channel Excel。
EXPORT_SITE_EXCELS = True

# 是否保留中间 memmap 文件。一般不需要。
KEEP_TEMP_MEMMAP = False

# 每处理多少个文件打印一次进度。
PRINT_EVERY_N_FILES = 1000

# header 固定 64 个 float32
HEADER_FLOATS = 64
HEADER_BYTES = HEADER_FLOATS * 4

# 文件名格式
FILENAME_RE = re.compile(
    r"(?P<Y>\d{4})-(?P<m>\d{2})-(?P<d>\d{2})-"
    r"(?P<H>\d{2})-(?P<M>\d{2})-(?P<S>\d{2})-out\.dat$",
    re.IGNORECASE
)


# =============================================================================
# 2. 基础函数
# =============================================================================

def choose_root_folder() -> Path:
    """UI 选择包含 2025-06-16、2025-06-17 ... 日期文件夹的总目录。"""
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    folder = filedialog.askdirectory(
        title="选择包含各日期 DAS 文件夹的总目录"
    )
    root.destroy()

    if not folder:
        print("未选择文件夹，程序结束。")
        sys.exit(0)

    return Path(folder).resolve()


def parse_file_time(path: Path) -> datetime | None:
    """从 2025-06-17-00-00-33-out.dat 解析时间。"""
    m = FILENAME_RE.match(path.name)
    if not m:
        return None
    try:
        return datetime(
            int(m.group("Y")),
            int(m.group("m")),
            int(m.group("d")),
            int(m.group("H")),
            int(m.group("M")),
            int(m.group("S")),
        )
    except ValueError:
        return None


def date_iter(d0, d1):
    """日期闭区间迭代。"""
    d = d0
    while d <= d1:
        yield d
        d += timedelta(days=1)


def read_header(path: Path):
    """
    只读取前 64 个 float32 header。
    返回 fs, n_channels, length_sec。
    """
    with open(path, "rb") as f:
        header = np.fromfile(f, dtype=np.float32, count=HEADER_FLOATS)

    if header.size < HEADER_FLOATS:
        raise ValueError(f"header 不足 64 个 float32：{path}")

    fs = float(header[10])
    n_channels = int(round(float(header[16])))
    length_sec = float(header[17])

    if not np.isfinite(fs) or fs <= 0:
        raise ValueError(f"无效采样率 fs={fs}：{path}")
    if n_channels <= 0:
        raise ValueError(f"无效通道数 N={n_channels}：{path}")
    if not np.isfinite(length_sec) or length_sec <= 0:
        raise ValueError(f"无效文件时长 length={length_sec}：{path}")

    return fs, n_channels, length_sec


def get_payload_sample_count(path: Path, n_channels: int) -> int:
    """
    根据真实文件字节数推断每通道 payload 样本数，
    避免 header length 与实际文件存在轻微不一致。
    """
    n_payload_bytes = path.stat().st_size - HEADER_BYTES
    if n_payload_bytes <= 0:
        raise ValueError(f"文件没有 payload：{path}")

    n_float = n_payload_bytes // 4
    if n_float % n_channels != 0:
        raise ValueError(
            f"payload float32 数量不能被通道数整除：{path}\n"
            f"payload floats={n_float}, N={n_channels}"
        )
    return int(n_float // n_channels)


def build_selected_channels(n_channels: int):
    """
    根据 SITE_CHANNEL_RANGES 生成：
    - selected_channels: 排序后的全研究通道索引
    - site_channels: 每个位置的通道数组
    """
    site_channels = {}
    all_channels = []

    for site, (c0, c1) in SITE_CHANNEL_RANGES.items():
        if c0 < 0 or c1 < c0:
            raise ValueError(f"{site} 通道范围错误：{c0}-{c1}")
        if c1 >= n_channels:
            raise ValueError(
                f"{site} 需要通道 {c0}-{c1}，"
                f"但原始文件只有 {n_channels} 个通道（最大索引 {n_channels-1}）。"
            )
        arr = np.arange(c0, c1 + 1, dtype=np.int32)
        site_channels[site] = arr
        all_channels.extend(arr.tolist())

    selected_channels = np.array(sorted(set(all_channels)), dtype=np.int32)
    return selected_channels, site_channels


def output_directory(root: Path) -> Path:
    name = (
        "DAS_1min_RMS_"
        f"{START_TIME:%Y%m%d_%H%M%S}_to_{END_TIME:%Y%m%d_%H%M%S}"
    )
    out = root / name
    out.mkdir(parents=True, exist_ok=True)
    return out


def minute_index(dt: datetime) -> int:
    return int((dt - START_TIME).total_seconds() // 60)


def sample_index_at_or_after(dt: datetime, file_start: datetime, fs: float) -> int:
    """
    返回第一个 sample_time >= dt 的样本索引。
    使用 ceil，并加极小容差，避免浮点误差。
    """
    x = (dt - file_start).total_seconds() * fs
    return int(math.ceil(x - 1e-9))


# =============================================================================
# 3. 文件发现
# =============================================================================

def find_first_valid_file(root: Path) -> Path:
    """
    在涉及日期文件夹中找到第一个命名有效且与目标时间附近相关的 dat，
    用于读取 N / fs。
    """
    scan_start_date = START_TIME.date() - timedelta(days=1)
    scan_end_date = END_TIME.date()

    for d in date_iter(scan_start_date, scan_end_date):
        day_dir = root / d.strftime("%Y-%m-%d")
        if not day_dir.is_dir():
            continue

        with os.scandir(day_dir) as it:
            for entry in it:
                if not entry.is_file():
                    continue
                if not entry.name.lower().endswith(".dat"):
                    continue
                p = Path(entry.path)
                ts = parse_file_time(p)
                if ts is not None:
                    return p

    raise FileNotFoundError(
        "在目标日期附近没有找到符合 "
        "'YYYY-MM-DD-HH-MM-SS-out.dat' 命名的 .dat 文件。"
    )


def collect_relevant_files_for_day(day_dir: Path,
                                   pad_seconds: float = 10.0):
    """
    只收集该天可能与目标时间窗有交集的文件。
    每天单独收集和排序，避免一次把 5 天所有路径放入内存。
    """
    files = []
    if not day_dir.is_dir():
        return files

    lower = START_TIME - timedelta(seconds=pad_seconds)
    upper = END_TIME + timedelta(seconds=pad_seconds)

    with os.scandir(day_dir) as it:
        for entry in it:
            if not entry.is_file():
                continue
            if not entry.name.lower().endswith(".dat"):
                continue

            p = Path(entry.path)
            ts = parse_file_time(p)
            if ts is None:
                continue

            if lower <= ts < upper:
                files.append((ts, p))

    files.sort(key=lambda x: x[0])
    return files


# =============================================================================
# 4. 核心：逐文件流式累加 1-min sum / sumsq / count
# =============================================================================

def process_all_files(root: Path, outdir: Path):
    first_file = find_first_valid_file(root)
    fs0, n_channels0, length0 = read_header(first_file)

    print("=" * 78)
    print("Qinhuangdao DAS raw .dat -> 1-min RMS")
    print("=" * 78)
    print(f"总目录：{root}")
    print(f"时间范围：[ {START_TIME} , {END_TIME} )")
    print(f"时长：{(END_TIME - START_TIME).total_seconds()/3600:.2f} h")
    print(f"首个有效文件：{first_file.name}")
    print(f"header 识别：fs={fs0:g} Hz, N={n_channels0}, length={length0:g} s")

    selected_channels, site_channels = build_selected_channels(n_channels0)
    n_sel = len(selected_channels)

    n_minutes_float = (END_TIME - START_TIME).total_seconds() / 60.0
    if abs(n_minutes_float - round(n_minutes_float)) > 1e-9:
        raise ValueError("START_TIME 到 END_TIME 不是整数分钟。")
    n_minutes = int(round(n_minutes_float))

    print(f"输出分钟数：{n_minutes}")
    print(f"研究通道数：{n_sel}")
    for site, arr in site_channels.items():
        print(f"  {site}: Ch {arr[0]}-{arr[-1]} ({len(arr)} channels)")

    # 通道索引 -> selected_channels 中的列位置
    selected_pos = {int(ch): i for i, ch in enumerate(selected_channels)}

    # 用磁盘 memmap 作为累加器，RAM 基本不随 5 天数据量增长。
    sum_path = outdir / "_tmp_sum.float64.dat"
    sumsq_path = outdir / "_tmp_sumsq.float64.dat"
    count_path = outdir / "_tmp_count.int32.dat"

    sums = np.memmap(
        sum_path, dtype=np.float64, mode="w+",
        shape=(n_minutes, n_sel)
    )
    sumsq = np.memmap(
        sumsq_path, dtype=np.float64, mode="w+",
        shape=(n_minutes, n_sel)
    )
    counts = np.memmap(
        count_path, dtype=np.int32, mode="w+",
        shape=(n_minutes, n_sel)
    )
    sums[:] = 0.0
    sumsq[:] = 0.0
    counts[:] = 0

    # 每分钟处理到的文件段数量，仅用于 QC。
    minute_segments = np.zeros(n_minutes, dtype=np.int32)

    processed_files = 0
    skipped_files = 0
    malformed_files = 0
    inconsistent_files = 0
    bytes_touched_est = 0
    t_start = time.perf_counter()

    # 扫描 start 前一天到 end 当天，防止恰好跨午夜的 5 s 文件遗漏。
    scan_start_date = START_TIME.date() - timedelta(days=1)
    scan_end_date = END_TIME.date()

    for d in date_iter(scan_start_date, scan_end_date):
        day_dir = root / d.strftime("%Y-%m-%d")
        if not day_dir.exists():
            # 当前案例中 start 前一天不存在不影响。
            if START_TIME.date() <= d <= END_TIME.date():
                print(f"[警告] 未找到日期文件夹：{day_dir}")
            continue

        day_files = collect_relevant_files_for_day(day_dir)
        if not day_files:
            if START_TIME.date() <= d <= END_TIME.date():
                print(f"[警告] {day_dir.name} 没有找到相关 dat 文件。")
            continue

        print(f"\n[{day_dir.name}] 待检查文件：{len(day_files)}")

        for file_ts, path in day_files:
            try:
                fs, n_channels, length_sec = read_header(path)

                if n_channels != n_channels0 or abs(fs - fs0) > 1e-6:
                    inconsistent_files += 1
                    print(
                        f"[跳过] fs/N 与首文件不一致：{path.name} | "
                        f"fs={fs}, N={n_channels}"
                    )
                    continue

                n_samples_payload = get_payload_sample_count(path, n_channels)

                # 优先相信实际 payload。
                n_samples = n_samples_payload
                duration_sec = n_samples / fs

                if FILE_TIMESTAMP_IS_START:
                    file_start = file_ts
                    file_end = file_start + timedelta(seconds=duration_sec)
                else:
                    file_end = file_ts
                    file_start = file_end - timedelta(seconds=duration_sec)

                # 与目标时段无交集。
                if file_end <= START_TIME or file_start >= END_TIME:
                    skipped_files += 1
                    continue

                overlap_start = max(file_start, START_TIME)
                overlap_end = min(file_end, END_TIME)
                if overlap_start >= overlap_end:
                    skipped_files += 1
                    continue

                # 只 mmap payload，不读取整个 5 s 文件进 RAM。
                mm = np.memmap(
                    path,
                    dtype=np.float32,
                    mode="r",
                    offset=HEADER_BYTES,
                    shape=(n_channels, n_samples),
                    order="F",
                )

                cursor = overlap_start

                while cursor < overlap_end:
                    this_minute_start = cursor.replace(second=0, microsecond=0)
                    this_minute_end = this_minute_start + timedelta(minutes=1)
                    seg_end = min(overlap_end, this_minute_end)

                    m_idx = minute_index(this_minute_start)
                    if not (0 <= m_idx < n_minutes):
                        cursor = seg_end
                        continue

                    i0 = sample_index_at_or_after(cursor, file_start, fs)
                    i1 = sample_index_at_or_after(seg_end, file_start, fs)

                    i0 = max(0, min(i0, n_samples))
                    i1 = max(0, min(i1, n_samples))

                    if i1 > i0:
                        # advanced indexing 只把所需通道、当前几秒切片读入 RAM。
                        block = np.asarray(
                            mm[selected_channels, i0:i1],
                            dtype=np.float64
                        )

                        finite = np.isfinite(block)

                        if finite.all():
                            # 最快路径：正常 DAS 一般全部有限。
                            sums[m_idx, :] += block.sum(axis=1, dtype=np.float64)
                            sumsq[m_idx, :] += np.einsum(
                                "ij,ij->i", block, block, dtype=np.float64
                            )
                            counts[m_idx, :] += block.shape[1]
                        else:
                            # 少数异常值时逐通道忽略 NaN/Inf。
                            safe = np.where(finite, block, 0.0)
                            sums[m_idx, :] += safe.sum(axis=1, dtype=np.float64)
                            sumsq[m_idx, :] += np.einsum(
                                "ij,ij->i", safe, safe, dtype=np.float64
                            )
                            counts[m_idx, :] += finite.sum(axis=1).astype(np.int32)

                        minute_segments[m_idx] += 1
                        bytes_touched_est += block.nbytes

                        del block, finite

                    cursor = seg_end

                del mm
                processed_files += 1

                if processed_files % PRINT_EVERY_N_FILES == 0:
                    elapsed = time.perf_counter() - t_start
                    rate = processed_files / max(elapsed, 1e-9)
                    print(
                        f"  已处理 {processed_files:,} files | "
                        f"{rate:.1f} files/s | "
                        f"elapsed {elapsed/60:.1f} min"
                    )
                    sums.flush()
                    sumsq.flush()
                    counts.flush()
                    gc.collect()

            except Exception as exc:
                malformed_files += 1
                print(f"[错误文件] {path.name}: {exc}")

        # 一天结束时 flush，降低意外中断损失。
        sums.flush()
        sumsq.flush()
        counts.flush()
        gc.collect()

    elapsed = time.perf_counter() - t_start

    print("\n" + "=" * 78)
    print("原始文件累计完成")
    print("=" * 78)
    print(f"processed files   : {processed_files:,}")
    print(f"skipped files     : {skipped_files:,}")
    print(f"malformed files   : {malformed_files:,}")
    print(f"inconsistent files: {inconsistent_files:,}")
    print(f"elapsed           : {elapsed/60:.2f} min")

    # =========================================================================
    # 计算 demeaned RMS:
    # sqrt( mean(x^2) - mean(x)^2 )
    # =========================================================================
    rms_path = outdir / "_tmp_rms.float32.dat"
    rms = np.memmap(
        rms_path, dtype=np.float32, mode="w+",
        shape=(n_minutes, n_sel)
    )
    rms[:] = np.nan

    # 分块按分钟计算，避免一次复制大矩阵。
    rows_per_chunk = 240
    for r0 in range(0, n_minutes, rows_per_chunk):
        r1 = min(n_minutes, r0 + rows_per_chunk)

        c = np.asarray(counts[r0:r1, :], dtype=np.float64)
        s = np.asarray(sums[r0:r1, :], dtype=np.float64)
        ss = np.asarray(sumsq[r0:r1, :], dtype=np.float64)

        valid_count = c > 0
        mean = np.full_like(s, np.nan)
        mean2 = np.full_like(ss, np.nan)

        np.divide(s, c, out=mean, where=valid_count)
        np.divide(ss, c, out=mean2, where=valid_count)

        var = mean2 - mean * mean
        # 浮点误差可能产生极小负数
        var = np.where(var < 0, np.maximum(var, 0.0), var)
        out = np.sqrt(var)

        # 每个通道独立按 90% 覆盖率判断。
        expected_per_minute = fs0 * 60.0
        coverage = c / expected_per_minute
        out[coverage < MIN_MINUTE_COVERAGE] = np.nan

        rms[r0:r1, :] = out.astype(np.float32)

        del c, s, ss, valid_count, mean, mean2, var, out, coverage

    rms.flush()

    # =========================================================================
    # QC：分钟覆盖率
    # 正常情况下各通道 count 相同，因此用中位数做 minute-level coverage。
    # =========================================================================
    expected_per_minute = fs0 * 60.0
    median_count = np.median(np.asarray(counts), axis=1)
    coverage_fraction = median_count / expected_per_minute
    valid_minute = coverage_fraction >= MIN_MINUTE_COVERAGE

    minute_times = [
        START_TIME + timedelta(minutes=i)
        for i in range(n_minutes)
    ]

    # CSV 用标准库，避免额外依赖 pandas。
    qc_csv = outdir / "minute_coverage.csv"
    with open(qc_csv, "w", encoding="utf-8-sig", newline="") as f:
        f.write(
            "time,median_sample_count,expected_sample_count,"
            "coverage_fraction,segments,valid_minute\n"
        )
        for i, dt in enumerate(minute_times):
            f.write(
                f"{dt:%Y-%m-%d %H:%M:%S},"
                f"{median_count[i]:.0f},"
                f"{expected_per_minute:.0f},"
                f"{coverage_fraction[i]:.6f},"
                f"{minute_segments[i]},"
                f"{int(valid_minute[i])}\n"
            )

    # =========================================================================
    # Excel 输出
    # =========================================================================
    if EXPORT_SITE_EXCELS:
        print("\n开始写 Excel（write-only 低内存模式）...")

        for site, channels in site_channels.items():
            out_xlsx = outdir / f"{site}_1min_RMS.xlsx"
            write_site_excel(
                out_xlsx=out_xlsx,
                site=site,
                channels=channels,
                selected_pos=selected_pos,
                minute_times=minute_times,
                rms=rms,
                fs=fs0,
                n_channels=n_channels0,
            )
            print(f"  完成：{out_xlsx.name}")

    # =========================================================================
    # summary
    # =========================================================================
    summary = outdir / "processing_summary.txt"
    with open(summary, "w", encoding="utf-8") as f:
        f.write("Qinhuangdao DAS raw .dat -> 1-min RMS\n")
        f.write("=" * 72 + "\n")
        f.write(f"Root folder: {root}\n")
        f.write(f"Start: {START_TIME}\n")
        f.write(f"End (exclusive): {END_TIME}\n")
        f.write(
            f"Duration_hours: "
            f"{(END_TIME - START_TIME).total_seconds()/3600:.3f}\n"
        )
        f.write(f"Number_of_minutes: {n_minutes}\n")
        f.write(f"Sampling_frequency_Hz: {fs0}\n")
        f.write(f"Raw_channel_count: {n_channels0}\n")
        f.write(f"Selected_channel_count: {n_sel}\n")
        f.write(f"Minimum_minute_coverage: {MIN_MINUTE_COVERAGE}\n")
        f.write(f"Processed_files: {processed_files}\n")
        f.write(f"Skipped_files: {skipped_files}\n")
        f.write(f"Malformed_files: {malformed_files}\n")
        f.write(f"Inconsistent_files: {inconsistent_files}\n")
        f.write(f"Valid_minutes: {int(valid_minute.sum())}/{n_minutes}\n")
        f.write(
            f"Median_coverage_fraction: "
            f"{np.nanmedian(coverage_fraction):.6f}\n"
        )
        f.write(
            f"Minimum_coverage_fraction: "
            f"{np.nanmin(coverage_fraction):.6f}\n"
        )
        f.write(f"Processing_elapsed_minutes: {elapsed/60:.3f}\n")
        f.write("\nRMS definition:\n")
        f.write(
            "RMS_1min = sqrt(mean(x^2) - mean(x)^2), "
            "computed from finite raw samples within each minute.\n"
        )
        f.write(
            "No log10 / winsorization / detrend / normalization / "
            "spectral analysis is applied at this stage.\n"
        )
        f.write("\nSite channel ranges:\n")
        for site, arr in site_channels.items():
            f.write(f"  {site}: {int(arr[0])}-{int(arr[-1])}\n")

    # 关闭 memmap 句柄
    sums.flush()
    sumsq.flush()
    counts.flush()
    rms.flush()

    del sums, sumsq, counts
    gc.collect()

    # rms 仍要在 Excel 完成后删除
    del rms
    gc.collect()

    # 删除中间文件
    if not KEEP_TEMP_MEMMAP:
        for p in [sum_path, sumsq_path, count_path, rms_path]:
            try:
                p.unlink(missing_ok=True)
            except Exception as exc:
                print(f"[提示] 临时文件未删除：{p.name} | {exc}")

    print("\n" + "=" * 78)
    print("全部完成")
    print("=" * 78)
    print(f"输出文件夹：{outdir}")
    print(f"分钟 QC：{qc_csv.name}")
    print(f"处理说明：{summary.name}")
    print(
        f"有效分钟：{int(valid_minute.sum())}/{n_minutes} "
        f"({100*valid_minute.mean():.2f}%)"
    )

    return outdir


# =============================================================================
# 5. 低内存 Excel 写出
# =============================================================================

def write_site_excel(out_xlsx: Path,
                     site: str,
                     channels: np.ndarray,
                     selected_pos: Dict[int, int],
                     minute_times: List[datetime],
                     rms: np.memmap,
                     fs: float,
                     n_channels: int):

    wb = Workbook(write_only=True)

    # ---------------- RMS sheet ----------------
    ws = wb.create_sheet("RMS_1min")

    header = ["channel"] + minute_times
    ws.append(header)

    for ch in channels:
        col = selected_pos[int(ch)]
        values = np.asarray(rms[:, col], dtype=np.float32)

        # openpyxl 对 np.float32 兼容并不总是稳定，转 Python float。
        row = [int(ch)]
        row.extend(
            None if not np.isfinite(v) else float(v)
            for v in values
        )
        ws.append(row)

    # ---------------- Metadata sheet ----------------
    meta = wb.create_sheet("Metadata")
    metadata_rows = [
        ("site", site),
        ("start_time", START_TIME),
        ("end_time_exclusive", END_TIME),
        ("minute_count", len(minute_times)),
        ("sampling_frequency_Hz", float(fs)),
        ("raw_channel_count", int(n_channels)),
        ("site_channel_start", int(channels[0])),
        ("site_channel_end", int(channels[-1])),
        ("minimum_minute_coverage", float(MIN_MINUTE_COVERAGE)),
        (
            "rms_definition",
            "sqrt(mean(x^2) - mean(x)^2) within each 1-min bin"
        ),
        (
            "processing_note",
            "raw 200-Hz -> 1-min demeaned RMS only; "
            "no log10/winsor/detrend/normalization"
        ),
    ]
    for row in metadata_rows:
        meta.append(row)

    wb.save(out_xlsx)


# =============================================================================
# 6. 主程序
# =============================================================================

def main():
    root = choose_root_folder()
    outdir = output_directory(root)

    try:
        result = process_all_files(root, outdir)

        # 完成后给一个轻量 UI 提示
        app = tk.Tk()
        app.withdraw()
        messagebox.showinfo(
            "DAS 1-min RMS 完成",
            f"处理完成。\n\n输出文件夹：\n{result}"
        )
        app.destroy()

    except Exception as exc:
        print("\n处理失败：")
        print(repr(exc))

        try:
            app = tk.Tk()
            app.withdraw()
            messagebox.showerror(
                "处理失败",
                f"{type(exc).__name__}: {exc}"
            )
            app.destroy()
        except Exception:
            pass

        raise


if __name__ == "__main__":
    main()
