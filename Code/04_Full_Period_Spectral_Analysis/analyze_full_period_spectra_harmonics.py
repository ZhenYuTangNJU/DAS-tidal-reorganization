from pathlib import Path
import warnings

import numpy as np
import pandas as pd
from scipy.signal import lombscargle
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2] / "03_Inputs_and_Intermediate_Results" / "Full_Period"
DATA_DIR = ROOT / "DAS_Merged"
TIDE_FILE = ROOT / "Tide_all.xlsx"
OUT = Path(__file__).resolve().parents[2] / "Reprocessed_Outputs" / "Full_Period_Spectral_Analysis"

SITE_FILES = {
    "Coast": DATA_DIR / "coast_merged.xlsx",
    "Western lagoon": DATA_DIR / "lagoon_1_merged.xlsx",
    "Northern lagoon": DATA_DIR / "lagoon_2_merged.xlsx",
    "Eastern lagoon": DATA_DIR / "lagoon_3_merged.xlsx",
}
SIGNAL_ORDER = ["Tide", *SITE_FILES.keys()]

# Frequencies in cycles per mean solar day. K2 is deliberately excluded because
# a 49-day record cannot resolve it from S2 (beat period about 182 days).
CONSTITUENTS = {
    "O1": 0.9295357,
    "K1": 1.0027379,
    "N2": 1.8959819,
    "M2": 1.9322736,
    "S2": 2.0000000,
    "M4": 3.8645472,
}
ROLLING_CONSTITUENTS = {k: CONSTITUENTS[k] for k in ["O1", "K1", "M2", "S2"]}



def read_das(path):
    raw = pd.read_excel(path, header=None)
    times = pd.to_datetime(raw.iloc[0, 1:], errors="coerce")
    keep_t = times.notna().to_numpy()
    times = pd.DatetimeIndex(times[keep_t])
    channels = pd.to_numeric(raw.iloc[1:, 0], errors="coerce").to_numpy(int)
    values = raw.iloc[1:, 1:].apply(pd.to_numeric, errors="coerce").to_numpy(float)[:, keep_t]
    return times, channels, values


def read_tide(path, start, end):
    raw = pd.read_excel(path, header=None)
    times = pd.to_datetime(raw.iloc[1:, 0], errors="coerce")
    values = pd.to_numeric(raw.iloc[1:, 1], errors="coerce")
    keep = times.notna() & values.notna() & (times >= start.floor("h")) & (times <= end.ceil("h"))
    return pd.DatetimeIndex(times[keep]), values[keep].to_numpy(float)


def preprocess_rows(times, values, use_log=True):
    """Winsorize, linearly detrend, and standardize each row without filling gaps."""
    t = (times - times[0]).total_seconds().to_numpy() / 86400.0
    out = np.full(values.shape, np.nan, float)
    for i, row in enumerate(values):
        y = row.astype(float).copy()
        if use_log:
            y[y <= 0] = np.nan
            y = np.log10(y)
        m = np.isfinite(y)
        if m.sum() < 10:
            continue
        qlo, qhi = np.nanpercentile(y[m], [0.5, 99.5])
        y[m] = np.clip(y[m], qlo, qhi)
        coef = np.polyfit(t[m], y[m], 1)
        y[m] -= np.polyval(coef, t[m])
        sd = np.nanstd(y[m], ddof=1)
        if sd > 0:
            y[m] /= sd
        out[i] = y
    return out


def preprocess_vector(times, values):
    t = (times - times[0]).total_seconds().to_numpy() / 86400.0
    y = values.astype(float).copy()
    m = np.isfinite(y)
    qlo, qhi = np.nanpercentile(y[m], [0.5, 99.5])
    y[m] = np.clip(y[m], qlo, qhi)
    coef = np.polyfit(t[m], y[m], 1)
    y[m] -= np.polyval(coef, t[m])
    y[m] /= np.nanstd(y[m], ddof=1)
    return y


def gls_spectrum(times, values, freq):
    m = np.isfinite(values)
    t = (times[m] - times[m][0]).total_seconds().to_numpy() / 86400.0
    y = values[m]
    p = lombscargle(t, y, 2 * np.pi * freq, normalize="power", floating_mean=True)
    return p


def harmonic_fit(times, values, constituents):
    m = np.isfinite(values)
    if m.sum() < 20:
        return None
    t = (times[m] - times[m][0]).total_seconds().to_numpy() / 86400.0
    y = values[m]
    cols = [np.ones_like(t), t - np.mean(t)]
    for f in constituents.values():
        cols.extend([np.cos(2 * np.pi * f * t), np.sin(2 * np.pi * f * t)])
    X = np.column_stack(cols)
    beta, _, _, _ = np.linalg.lstsq(X, y, rcond=None)
    fit = X @ beta
    resid = y - fit
    ss_tot = np.sum((y - np.mean(y)) ** 2)
    r2 = 1 - np.sum(resid ** 2) / ss_tot if ss_tot > 0 else np.nan
    amps, phases = {}, {}
    for j, name in enumerate(constituents):
        a, b = beta[2 + 2*j], beta[3 + 2*j]
        amps[name] = float(np.hypot(a, b))
        phases[name] = float(np.degrees(np.arctan2(-b, a)) % 360)
    powers = np.array([amps[k] ** 2 for k in constituents])
    fractions = powers / powers.sum() if powers.sum() > 0 else np.full_like(powers, np.nan)
    return {
        "n": int(m.sum()), "span_days": float(t.max() - t.min()), "r2": float(r2),
        "condition_number": float(np.linalg.cond(X)),
        "amplitude": amps, "phase_deg": phases,
        "fraction": dict(zip(constituents.keys(), fractions)),
    }












def format_workbook(path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for col in range(1, ws.max_column + 1):
            vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 100) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = min(34, max(10, max(map(len, vals)) + 2))
    wb.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    das = {}
    audit = []
    reference_times = None
    for site, path in SITE_FILES.items():
        times, channels, values = read_das(path)
        if reference_times is None:
            reference_times = times
        elif not np.array_equal(reference_times.values, times.values):
            raise ValueError("DAS time axes differ among sector files")
        z = preprocess_rows(times, values, use_log=True)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", category=RuntimeWarning)
            median = np.nanmedian(z, axis=0)
        das[site] = {"times": times, "channels": channels, "raw": values, "z": z, "median": median}
        mask = np.isfinite(values[0])
        audit.append({
            "record": site, "start": times.min(), "end": times.max(), "calendar_span_days": (times.max()-times.min()).total_seconds()/86400,
            "cadence_minutes": np.median(np.diff(times.values).astype("timedelta64[m]").astype(float)),
            "n_channels": len(channels), "n_time_slots": len(times), "n_valid_slots": int(mask.sum()),
            "temporal_coverage": float(mask.mean()), "all_channels_share_time_mask": True,
            "preprocessing": "log10; 0.5-99.5% winsorize; linear detrend; within-channel z-score",
        })
    start, end = reference_times.min(), reference_times.max()
    tide_times, tide_raw = read_tide(TIDE_FILE, start, end)
    tide_z = preprocess_vector(tide_times, tide_raw)
    audit.append({
        "record": "Tide", "start": tide_times.min(), "end": tide_times.max(), "calendar_span_days": (tide_times.max()-tide_times.min()).total_seconds()/86400,
        "cadence_minutes": np.median(np.diff(tide_times.values).astype("timedelta64[m]").astype(float)),
        "n_channels": 1, "n_time_slots": len(tide_times), "n_valid_slots": int(np.isfinite(tide_raw).sum()),
        "temporal_coverage": float(np.isfinite(tide_raw).mean()), "all_channels_share_time_mask": True,
        "preprocessing": "0.5-99.5% winsorize; linear detrend; z-score; restricted to DAS calendar span",
    })

    freq = np.linspace(0.2, 4.3, 4101)
    spectra = {"Tide": gls_spectrum(tide_times, tide_z, freq)}
    for site in SITE_FILES:
        spectra[site] = gls_spectrum(das[site]["times"], das[site]["median"], freq)
    das_mask = np.isfinite(das["Coast"]["median"])

    # Full-record fixed-frequency harmonic regression.
    tide_fit = harmonic_fit(tide_times, tide_z, CONSTITUENTS)
    harmonic_signal_rows = []
    for name, fit in [("Tide", tide_fit)] + [(s, harmonic_fit(das[s]["times"], das[s]["median"], CONSTITUENTS)) for s in SITE_FILES]:
        row = {"signal": name, "n": fit["n"], "span_days": fit["span_days"], "R2": fit["r2"], "condition_number": fit["condition_number"]}
        for k in CONSTITUENTS:
            row[f"{k}_amplitude_z"] = fit["amplitude"][k]
            row[f"{k}_phase_deg_diagnostic"] = fit["phase_deg"][k]
            row[f"{k}_power_fraction"] = fit["fraction"][k]
        harmonic_signal_rows.append(row)

    tide_fraction = np.array([tide_fit["fraction"][k] for k in CONSTITUENTS])
    channel_rows = []
    for site in SITE_FILES:
        for ch, y in zip(das[site]["channels"], das[site]["z"]):
            fit = harmonic_fit(das[site]["times"], y, CONSTITUENTS)
            comp = np.array([fit["fraction"][k] for k in CONSTITUENTS])
            row = {"site": site, "channel": int(ch), "n": fit["n"], "R2": fit["r2"],
                   "condition_number": fit["condition_number"], "THI6": 0.5*np.sum(np.abs(comp-tide_fraction))}
            for k in CONSTITUENTS:
                row[f"{k}_power_fraction"] = fit["fraction"][k]
                row[f"{k}_amplitude_z"] = fit["amplitude"][k]
                row[f"{k}_phase_deg_diagnostic"] = fit["phase_deg"][k]
            channel_rows.append(row)
    channel_metrics = pd.DataFrame(channel_rows)

    # 28-day moving windows, 3-day step. The reduced four-constituent set avoids
    # unstable N2-M2 separation in moving windows.
    window_days, step_days = 28, 3
    first_center = start + pd.Timedelta(days=window_days/2)
    last_center = end - pd.Timedelta(days=window_days/2)
    centers = pd.date_range(first_center.floor("D"), last_center.ceil("D"), freq=f"{step_days}D")
    rolling_rows = []
    for center in centers:
        lo, hi = center - pd.Timedelta(days=window_days/2), center + pd.Timedelta(days=window_days/2)
        tm = (tide_times >= lo) & (tide_times <= hi)
        tide_win = harmonic_fit(tide_times[tm], tide_z[tm], ROLLING_CONSTITUENTS)
        if tide_win is None:
            continue
        tide_comp4 = np.array([tide_win["fraction"][k] for k in ROLLING_CONSTITUENTS])
        for site in SITE_FILES:
            times = das[site]["times"]
            m = (times >= lo) & (times <= hi)
            fit = harmonic_fit(times[m], das[site]["median"][m], ROLLING_CONSTITUENTS)
            if fit is None:
                continue
            comp = np.array([fit["fraction"][k] for k in ROLLING_CONSTITUENTS])
            row = {"window_center": center, "window_start": lo, "window_end": hi, "window_days": window_days,
                   "signal": site, "n_valid": fit["n"], "R2": fit["r2"], "condition_number": fit["condition_number"],
                   "THI4_vs_tide": 0.5*np.sum(np.abs(comp-tide_comp4))}
            for k in ROLLING_CONSTITUENTS:
                row[f"{k}_power_fraction"] = fit["fraction"][k]
                row[f"tide_{k}_power_fraction"] = tide_win["fraction"][k]
            rolling_rows.append(row)
    rolling = pd.DataFrame(rolling_rows)

    # Peak summary uses the local maximum within +/-0.025 cpd of each nominal frequency.
    peak_rows = []
    for signal in SIGNAL_ORDER:
        p = spectra[signal]
        for name, f0 in CONSTITUENTS.items():
            m = np.abs(freq-f0) <= 0.025
            idx_local = np.argmax(p[m])
            inds = np.where(m)[0]
            idx = inds[idx_local]
            peak_rows.append({"signal": signal, "constituent": name, "nominal_frequency_cpd": f0,
                              "local_peak_frequency_cpd": freq[idx], "local_peak_power": p[idx],
                              "power_normalized_to_signal_max": p[idx]/np.max(p)})

    spec_df = pd.DataFrame({"frequency_cpd": freq, **{s: spectra[s] for s in SIGNAL_ORDER}})
    audit_df = pd.DataFrame(audit)
    harmonic_signal_df = pd.DataFrame(harmonic_signal_rows)
    peak_df = pd.DataFrame(peak_rows)
    definitions = pd.DataFrame([
        ["GLS spectrum", "Generalized Lomb-Scargle periodogram with floating mean on observed samples only; no gap filling."],
        ["Harmonic power fraction", "Squared fitted amplitude for one constituent divided by the sum across the stated constituent set."],
        ["THI6", "0.5 times the L1 distance between a channel's six-constituent power-fraction vector and the tide vector."],
        ["THI4_vs_tide", "Moving-window equivalent using O1, K1, M2, and S2 in a 28-day window."],
        ["Phase columns", "Diagnostic regression phases only; excluded from inferential conclusions because DAS intensity is not water level and clock/response lags may contribute."],
    ], columns=["metric", "definition"])

    xlsx = OUT / "full_period_spectral_harmonic_source_data.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        audit_df.to_excel(writer, sheet_name="data_audit", index=False)
        spec_df.to_excel(writer, sheet_name="GLS_spectra", index=False)
        peak_df.to_excel(writer, sheet_name="spectral_peaks", index=False)
        harmonic_signal_df.to_excel(writer, sheet_name="full_harmonic_signals", index=False)
        channel_metrics.to_excel(writer, sheet_name="channel_harmonics", index=False)
        rolling.to_excel(writer, sheet_name="rolling_28d_harmonics", index=False)
        definitions.to_excel(writer, sheet_name="metric_definitions", index=False)
    format_workbook(xlsx)

    # Compact machine-readable summaries for audit and manuscript drafting.
    summary = channel_metrics.groupby("site").agg(
        n_channels=("channel", "size"), THI6_median=("THI6", "median"),
        THI6_q25=("THI6", lambda x: np.quantile(x, 0.25)), THI6_q75=("THI6", lambda x: np.quantile(x, 0.75)),
        harmonic_R2_median=("R2", "median"), harmonic_R2_q25=("R2", lambda x: np.quantile(x, 0.25)),
        harmonic_R2_q75=("R2", lambda x: np.quantile(x, 0.75)),
    ).reset_index()
    summary.to_csv(OUT / "sector_harmonic_summary.csv", index=False, encoding="utf-8-sig")

    print("OUTPUT", OUT)
    print(audit_df[["record", "calendar_span_days", "n_valid_slots", "temporal_coverage"]].to_string(index=False))
    print("\nSECTOR SUMMARY")
    print(summary.to_string(index=False))
    print("\nSIGNAL COMPOSITIONS")
    print(harmonic_signal_df[["signal", "R2"] + [f"{k}_power_fraction" for k in CONSTITUENTS]].to_string(index=False))
    print("\nROLLING CONDITION NUMBER MAX", rolling["condition_number"].max())


if __name__ == "__main__":
    main()
